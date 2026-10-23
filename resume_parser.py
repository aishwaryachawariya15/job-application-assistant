"""Resume parsing: extract raw text from PDF / DOCX / TXT resumes and
derive a ranked list of keywords/skills used for job matching.
"""
from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field

from skills import STOPWORDS, count_skill_occurrences, find_skills_in_text

WORD_RE = re.compile(r"[A-Za-z][A-Za-z+#.\-]*")


@dataclass
class ParsedResume:
    name: str
    raw_text: str
    keywords: list[str] = field(default_factory=list)
    matched_skills: list[str] = field(default_factory=list)
    word_count: int = 0
    sections_found: list[str] = field(default_factory=list)
    has_contact_info: bool = False


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from a resume file based on its extension."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext in ("txt", "md"):
        return file_bytes.decode("utf-8", errors="ignore")

    # Fallback: try to decode as text.
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError(
            "pdfplumber is required to read PDF resumes. Install with: pip install pdfplumber"
        ) from exc

    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required to read DOCX resumes. Install with: pip install python-docx"
        ) from exc

    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    # Also pull text from tables, since many resumes use table layouts.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _ngrams(words: list[str], n: int) -> list[str]:
    return [" ".join(words[i : i + n]) for i in range(len(words) - n + 1)]


def extract_keywords(text: str, top_n: int = 25) -> tuple[list[str], list[str]]:
    """Return (ranked_keywords, matched_known_skills) from resume text.

    Ranked keywords combine frequency-based unigrams/bigrams with any
    recognized skills from SKILLS_DB that appear in the text.
    """
    lower_text = text.lower()

    matched_skills = sorted(
        find_skills_in_text(lower_text),
        key=lambda s: count_skill_occurrences(lower_text, s),
        reverse=True,
    )

    words = [w.lower() for w in WORD_RE.findall(text)]
    filtered = [w for w in words if w not in STOPWORDS and len(w) > 2]

    unigram_counts = Counter(filtered)
    bigram_counts = Counter(_ngrams(filtered, 2))

    # Drop bigrams that are just stopword combinations / noise, keep ones
    # that look like meaningful phrases (both words length > 2 already).
    ranked_bigrams = [phrase for phrase, _ in bigram_counts.most_common(30)]
    ranked_unigrams = [word for word, _ in unigram_counts.most_common(60)]

    combined: list[str] = []
    # Known skills first (highest confidence signal).
    combined.extend(matched_skills)
    for phrase in ranked_bigrams:
        if phrase not in combined:
            combined.append(phrase)
    for word in ranked_unigrams:
        if word not in combined:
            combined.append(word)

    return combined[:top_n], matched_skills


def detect_sections(text: str) -> list[str]:
    lower_text = text.lower()
    section_markers = [
        "experience", "work experience", "professional experience",
        "education", "skills", "summary", "objective", "certifications",
        "projects", "achievements",
    ]
    return [s for s in section_markers if s in lower_text]


def detect_contact_info(text: str) -> bool:
    has_email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text) is not None
    has_phone = re.search(r"(\+?\d[\d\-\s().]{8,}\d)", text) is not None
    return has_email or has_phone


def parse_resume(filename: str, file_bytes: bytes) -> ParsedResume:
    raw_text = extract_text(filename, file_bytes)
    keywords, matched_skills = extract_keywords(raw_text)
    sections = detect_sections(raw_text)
    has_contact = detect_contact_info(raw_text)
    word_count = len(WORD_RE.findall(raw_text))

    return ParsedResume(
        name=filename,
        raw_text=raw_text,
        keywords=keywords,
        matched_skills=matched_skills,
        word_count=word_count,
        sections_found=sections,
        has_contact_info=has_contact,
    )
